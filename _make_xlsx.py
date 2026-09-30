# -*- coding: utf-8 -*-
"""
Generate 全站文案总表-europastudy.com.xlsx (master copy table).
Structure:
  Sheet 0: 使用说明
  Sheet 1: 全站通用(导航 + 两套页脚, 标注适用页面)
  Sheet 2-9: 每个页面一张 (只含该页正文区块; 导航/页脚已在通用表)
Columns: 序号 | 网页位置 | 现在的英文文案 | 想改成什么 | 备注 | 内部定位(勿改)
"""
import os, re
from bs4 import BeautifulSoup, NavigableString, Tag
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = r"E:/WORKBUDDY/2026-09-03-16-18-34/company-website-preview"
OUT = r"D:/我的坚果云/我的坚果云/网站/中西桥教育英文官网项目/02-内容清单/全站文案总表-europastudy.com.xlsx"

PAGES = [
    ("index.html",                 "① 首页 Home"),
    ("about.html",                 "② 关于我们 About us"),
    ("partners.html",              "③ 合作大学 Partners"),
    ("cooperative-education.html", "④ 中外合作办学 Cooperative education"),
    ("tailor-made-master.html",    "⑤ 硕士专班 Exclusive master's"),
    ("admissions.html",            "⑥ 申请与录取 Admissions"),
    ("study-visits.html",          "⑦ 研学与短期项目 Study visits"),
    ("contact.html",               "⑧ 联系我们 Contact"),
]

BLOCK = {"h1","h2","h3","h4","h5","h6","p","li","a","button","figcaption","td","th",
         "address","blockquote","caption","dt","dd","label"}
LAYOUT = {"div","section","footer","header","nav","ul","ol","figure","main","form","aside"}
INLINE = {"span","b","strong","em","i","small","sub","sup","code","br","abbr","u"}
STRIP  = {"script","style","noscript","svg","template"}
SEM = ["svc-card","uni-card","uni-tile","why-card","card","metric","tl-item","step",
       "office","feat-card","visit","car-slide","prose"]
SEM_ZH = {"svc-card":"服务卡","uni-card":"大学卡","uni-tile":"大学卡","why-card":"优势卡",
          "card":"卡片","metric":"数据条","tl-item":"时间线条目","step":"步骤",
          "office":"办公室","feat-card":"亮点卡","visit":"访问图","car-slide":"轮播图","prose":"正文块"}

# ---------- styles ----------
NAVY   = "1F3864"; BLUE = "2E5B9F"; LIGHT = "DCE6F1"; YELLOW = "FFF2CC"; GREY = "808080"
HDR_FILL = PatternFill("solid", fgColor=NAVY)
SUB_FILL = PatternFill("solid", fgColor=BLUE)
LIGHT_FILL = PatternFill("solid", fgColor=LIGHT)
YELL_FILL = PatternFill("solid", fgColor=YELLOW)
HDR_FONT  = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
SUB_FONT  = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
SEC_FONT  = Font(name="Microsoft YaHei", size=10, bold=True, color=NAVY)
NOTE_FONT = Font(name="Microsoft YaHei", size=9, color=GREY, italic=True)
BODY_FONT = Font(name="Microsoft YaHei", size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
WRAP_C = Alignment(wrap_text=True, vertical="center", horizontal="center")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def clean(t):
    t = t.replace("\u200b","").replace("\ufeff","")
    return re.sub(r"\s+", " ", t).strip()

def has_block_desc(el):
    return el.find(BLOCK) is not None

def inline_segments(cont):
    """Return [(el_or_None, text)] for direct inline children / direct text of a layout container."""
    segs = []
    for child in cont.children:
        if isinstance(child, NavigableString):
            t = clean(str(child))
            if t: segs.append((None, t))
        elif isinstance(child, Tag) and child.name in INLINE:
            t = clean(child.get_text(" "))
            if t: segs.append((child, t))
    return segs

def row_parts(cont):
    """Return list of (el_or_None, text) rows. BLOCK -> whole text with self; layout container -> per inline child."""
    if cont.name == "figcaption":
        # figcaption often mixes <b> title with bare text caption
        segs = []
        for child in cont.children:
            if isinstance(child, NavigableString):
                t = clean(str(child))
                if t: segs.append((None, t))
            elif isinstance(child, Tag) and child.name in INLINE:
                t = clean(child.get_text(" "))
                if t: segs.append((child, t))
        if len(segs) >= 2:
            return segs
        return segs if segs else [(cont, clean(cont.get_text(" ")))]
    if cont.name in BLOCK:
        return [(cont, clean(cont.get_text(" ")))]
    segs = inline_segments(cont)
    if len(segs) <= 1:
        return segs if segs else [(cont, clean(cont.get_text(" ")))]
    return segs

def iter_rows(container):
    """Yield (container_el, [(sub_el_or_None, text), ...])."""
    seen = set()
    for tn in container.find_all(string=True):
        if not str(tn).strip():
            continue
        parent = tn.parent
        if parent.name in STRIP:
            continue
        cont = None
        cur = parent
        while cur is not None and cur.name not in ("html","body"):
            if cur.name in BLOCK:
                cont = cur; break
            if cur.name in LAYOUT and not has_block_desc(cur):
                cont = cur; break
            cur = cur.parent
        if cont is None:
            # fallback: bare inline holder (span.badge etc.) whose parent is layout
            if parent.name in INLINE and parent.parent is not None and parent.parent.name in LAYOUT:
                cont = parent
            else:
                continue
        if id(cont) in seen:
            continue
        pieces = row_parts(cont)
        pieces = [(e, t) for e, t in pieces if t]
        if not pieces:
            seen.add(id(cont)); continue
        joint = " | ".join(t for _, t in pieces)
        stripped = re.sub(r"[\s\|\u00b7\u2022\u2013\u2014_/\\*·•✦→←‹›«»▾✓☑]", "", joint)
        if not stripped:
            seen.add(id(cont)); continue
        cls = cont.get("class") or []
        if re.fullmatch(r"0?\d{1,3}", stripped) and cont.name != "h3" and any("num" in c or c=="ic" for c in cls):
            seen.add(id(cont)); continue
        # skip decorative emoji/icon glyphs (e.g. inside .ic, or lone symbols)
        if any("ic" in c for c in cls) and not re.search(r"[A-Za-z0-9\u00c0-\u024f]", joint):
            seen.add(id(cont)); continue
        seen.add(id(cont))
        yield cont, pieces

def piece_label(cont, sub, idx):
    """Refine row label for pieces inside composite containers (metrics / figcaption)."""
    cls = cont.get("class") or []
    if any("metric" in c for c in cls) and cont.name != "span":
        return "数字/年份" if idx == 0 else "说明文字"
    if cont.name == "figcaption":
        return "标题行(粗体)" if idx == 0 else "说明行"
    if sub is not None:
        return elem_label(sub)
    return elem_label(cont)

def sem_card(cont):
    """Find nearest semantic card ancestor -> (zh label, order index)."""
    cur = cont
    while cur is not None and cur.name not in ("html","body"):
        cls = cur.get("class") or []
        for c in cls:
            if c in SEM:
                # count how many elements with that class precede it
                return c, cls2idx[c].get(id(cur), 0)
        cur = cur.parent
    return None, None

cls2idx = {}

def prepare_counters(soup):
    for c in SEM:
        els = soup.find_all(class_=c)
        cls2idx[c] = {id(e): i+1 for i, e in enumerate(els)}

ELEM_META = [
    ("p.lead","导语 / 副标题"), ("span.badge","徽标行"), ("h1","页面大标题"), ("h2","区块主标题"),
    ("h3","卡片标题"), ("h4","栏目标题"), ("h5","小标题"), ("p","段落文字"), ("li","要点列表项"),
    ("a.btn","按钮文字"), ("a","链接文字"), ("button","按钮"), ("figcaption","图片说明"),
    ("span","小字/标签"), ("strong","强调文字"), ("b","强调文字"), ("em","强调斜体"),
    ("td","表格文字"), ("th","表格表头"), ("div","文字块"),
]
CLASS_LABELS = {  # pure class-name semantics
    "eyebrow":"小眉题(区块小标签)","place":"城市/属性","role":"角色/职能说明","tag2":"标签文字",
    "svc-more":"“了解更多”链接","car-caption":"轮播图说明","ic":"图标(无文字)","lead":"导语",
}

def elem_label(cont):
    sel = cont.name + (("." + ".".join(cont.get("class") or [])) if cont.get("class") else "")
    cls = cont.get("class") or []
    # 1) named class lookups (eyebrow/place/role/tag2/svc-more/car-caption...)
    for c in cls:
        if c in CLASS_LABELS:
            return CLASS_LABELS[c]
    # 2) combined selector like p.lead, span.badge, a.btn
    for pat, zh in ELEM_META:
        if sel == pat or (pat.startswith(cont.name + ".") and sel.startswith(pat)):
            return zh
    # 3) plain tag match
    for pat, zh in ELEM_META:
        if pat == cont.name and "." not in pat:
            return zh
    return "·".join(cls[:2]) if cls else cont.name

def css_of(cont):
    """Minimal css-ish locator for the internal column."""
    path = []
    cur = cont
    while cur is not None and cur.name not in ("html","body"):
        tag = cur.name
        if cur.get("id"):
            tag += "#" + cur["id"]
        cls = cur.get("class") or []
        if cls:
            tag += "." + ".".join(cls[:2])
        path.append(tag)
        cur = cur.parent
    return " > ".join(reversed(path))

def extract_page(fname):
    soup = BeautifulSoup(open(os.path.join(SRC, fname), encoding="utf-8").read(), "html.parser")
    for t in soup(STRIP):
        t.decompose()
    prepare_counters(soup)
    title = clean(soup.title.get_text()) if soup.title else ""
    body = soup.body or soup
    blocks = []
    # nav
    hdr = soup.find("header")
    if hdr:
        rows = []
        for cont, pieces in iter_rows(hdr):
            for idx, (sub, pc) in enumerate(pieces):
                txt = pc
                # nav trigger 'Services ▾' strip caret
                if txt.endswith("▾"):
                    txt = txt[:-1].strip()
                if not txt: continue
                rows.append({"el": piece_label(cont, sub, idx), "txt": txt, "css": css_of(cont)})
        blocks.append(("nav", "顶部导航(全站通用)", rows))
    # main/body sections (incl. non-<section> content blocks such as div.feat)
    container = soup.find("main") or body
    idx = 0
    for el in container.find_all(recursive=False):
        if el.name not in ("section","div"):
            continue
        if el.name == "div" and el.find(["h1","h2","h3","p","li","figcaption"]) is None:
            continue  # empty layout wrapper
        if "page-hero" in (el.get("class") or []):
            rows = []
            for cont, pieces in iter_rows(el):
                for idx2, (sub, pc) in enumerate(pieces):
                    rows.append({"el": piece_label(cont, sub, idx2), "txt": pc, "css": css_of(cont)})
            blocks.append(("hero","横幅区(页面头部)", rows))
            continue
        if el.name == "section":
            idx += 1
            h2 = el.find("h2")
            h1 = el.find("h1")
            ey = el.find("p.eyebrow")
            secid = el.get("id") or ""
            idzh = {"about":"关于我们(简介)","universities":"合作大学墙(6校logo)",
                    "services":"我们能做什么(4大服务)","why":"为什么选择我们",
                    "partners":"合作大学列表"}.get(secid, "")
            if h2 is not None:
                ttl = clean(h2.get_text(" "))
            elif ey is not None:
                ttl = clean(ey.get_text(" "))
            elif h1 is not None:
                ttl = "横幅区(首屏大标题区)"   # h1 itself is a row below
            elif idzh:
                ttl = idzh
            else:
                ttl = ""
            rows = []
            for cont, pieces in iter_rows(el):
                card, cnum = sem_card(cont)
                for idx2, (sub, pc) in enumerate(pieces):
                    lbl = piece_label(cont, sub, idx2)
                    if card and cnum:
                        lbl = f"{SEM_ZH.get(card,card)} {cnum} · {lbl}"
                    rows.append({"el": lbl, "txt": pc, "css": css_of(cont)})
            blocks.append(("sec", ttl or f"区块 {idx}", rows))
        else:
            # standalone content div (div.feat etc.)
            idx += 1
            h2 = el.find("h2")
            ey = el.find("p.eyebrow")
            h1 = el.find("h1")
            if h2 is not None:
                ttl = clean(h2.get_text(" "))
            elif ey is not None:
                ttl = clean(ey.get_text(" "))
            elif h1 is not None:
                ttl = "横幅区(首屏大标题区)"
            else:
                ttl = ""
            rows = []
            for cont, pieces in iter_rows(el):
                card, cnum = sem_card(cont)
                for idx2, (sub, pc) in enumerate(pieces):
                    lbl = piece_label(cont, sub, idx2)
                    if card and cnum:
                        lbl = f"{SEM_ZH.get(card,card)} {cnum} · {lbl}"
                    rows.append({"el": lbl, "txt": pc, "css": css_of(cont)})
            blocks.append(("sec", ttl or f"区块 {idx}", rows))
    # footer
    ft = soup.find("footer")
    if ft:
        rows = []
        for cont, pieces in iter_rows(ft):
            for idx2, (sub, pc) in enumerate(pieces):
                rows.append({"el": piece_label(cont, sub, idx2), "txt": pc, "css": css_of(cont)})
        blocks.append(("footer", "页脚(全站通用)", rows))
    return title, blocks

def styled_sheet(ws, title_row):
    """Header row style."""
    for c in range(1, 7):
        cell = ws.cell(row=title_row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = WRAP_C
        cell.border = BORDER

def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def write_block(ws, r, label, rows, prefix):
    """Write a block of rows starting at r. Returns next free row."""
    if not rows:
        return r
    # section label row
    ws.cell(row=r, column=1, value="")
    ws.cell(row=r, column=2, value=label)
    ws.cell(row=r, column=2).font = SEC_FONT
    ws.cell(row=r, column=2).fill = LIGHT_FILL
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    ws.cell(row=r, column=2).alignment = Alignment(vertical="center")
    ws.row_dimensions[r].height = 18
    r += 1
    for i, row in enumerate(rows, 1):
        ws.cell(row=r, column=1, value=i).alignment = WRAP_C
        ws.cell(row=r, column=2, value=row["el"])
        ws.cell(row=r, column=3, value=row["txt"])
        ws.cell(row=r, column=5, value=row.get("note",""))
        ws.cell(row=r, column=6, value=row["css"])
        for col in range(1, 7):
            cell = ws.cell(row=r, column=col)
            cell.border = BORDER
            cell.font = BODY_FONT
            if col in (2,3,5):
                cell.alignment = WRAP
        ws.cell(row=r, column=6).font = NOTE_FONT
        r += 1
    return r

# ================= build =================
wb = Workbook()

# ---------- Sheet 0 使用说明 ----------
ws = wb.active
ws.title = "0 使用说明"
set_widths(ws, [2, 18, 110])
ws.sheet_view.showGridLines = False
lines = [
    ("", "", ""),
    ("说明", "这份表是什么", "这是 europastudy.com 全站文案总表:把网站上每一处可见英文文字都列出来了,方便你自己核对和修改措辞。"),
    ("说明", "怎么改", "打开对应页面的工作表 → 在 C 列找到想改的句子 → 在 D 列【想改成什么】里直接写新的英文。不会英文也可以写中文,我会帮你翻成合适的英文再上线。"),
    ("说明", "不改怎么办", "D 列留空 = 这行不改。不需要整张表都填,只填你想动的行。"),
    ("说明", "改完怎么交给我", "保存文件,回到对话框把文件发我(或直接告诉我改到第几行/哪句话),我会同步到网站并发布,一般 10 分钟内生效。"),
    ("说明", "小心页脚两套", "网站目前有两套页脚(首页/关于我们/合作大学 是一套;其余页面是另一套)。如果你改页脚措辞,两套都要改,或者直接告诉我‘把页脚统一成一套’,我来处理。"),
    ("说明", "哪些别乱动", "机构全称、大学名称、地址、邮箱、电话、年份数字等事实信息建议保留;按钮/链接文字改了要注意和指向的页面一致。"),
    ("说明", "F 列是什么", "灰色 F 列是我定位代码用的内部标记,请勿删除或修改。"),
    ("说明", "页面结构", "①-⑧ 号工作表 = 8 个页面正文。顶部导航和页脚因为每页都有,单独放在【1 全站通用】里,在那里改一次=全部页面一起变。"),
]
r = 2
for kind, a, b in lines:
    if kind == "说明":
        ws.cell(row=r, column=2, value=a).font = Font(name="Microsoft YaHei", size=11, bold=True, color=NAVY)
        cell = ws.cell(row=r, column=3, value=b)
        cell.font = BODY_FONT; cell.alignment = WRAP
        ws.row_dimensions[r].height = max(18, 16 * (len(b)//70 + 1))
        r += 1
# title
ws.cell(row=1, column=1, value="全站文案总表 — europastudy.com(中西桥教育英文官网)")
ws.cell(row=1, column=1).font = Font(name="Microsoft YaHei", size=14, bold=True, color="FFFFFF")
ws.cell(row=1, column=1).fill = HDR_FILL
ws.merge_cells("A1:C1")
ws.row_dimensions[1].height = 26

# ---------- Sheet 1 全站通用 ----------
ws = wb.create_sheet("1 全站通用(导航·页脚)")
set_widths(ws, [6, 22, 52, 40, 26, 60])
headers = ["序号","网页位置","现在的英文文案","想改成什么(留空=不改)","备注","内部定位(勿改)"]
for i, h in enumerate(headers, 1):
    ws.cell(row=1, column=i, value=h)
styled_sheet(ws, 1)
ws.freeze_panes = "A2"

r = 2
# --- nav from index ---
soup = BeautifulSoup(open(os.path.join(SRC,"index.html"), encoding="utf-8").read(), "html.parser")
for t in soup(STRIP): t.decompose()
hdr = soup.find("header")
nav_rows = []
for cont, pieces in iter_rows(hdr):
    for idx, (sub, pc) in enumerate(pieces):
        txt = pc[:-1].strip() if pc.endswith("▾") else pc
        if not txt: continue
        note = "下拉菜单按钮" if "drop" in css_of(cont) and "drop-menu" not in css_of(cont) else ""
        nav_rows.append({"el": elem_label(sub or cont), "txt": txt, "css": css_of(cont), "note": note})
# manual nav structure presentation
nav_struct = [
    ("logo","Logo(图片,无文字)","","点击回首页"),
]
nav_rows2 = []
for row in nav_rows:
    css = row["css"]
    if "drop-menu" in css:
        nav_rows2.append({"el":"下拉菜单项","txt":row["txt"],"css":row["css"],"note":"Services 下拉里的一项"})
    elif "drop" in css:
        nav_rows2.append({"el":"导航按钮(带下拉)","txt":row["txt"],"css":row["css"],"note":"悬停/点击展开下拉菜单"})
    else:
        nav_rows2.append({"el":"导航链接","txt":row["txt"],"css":row["css"],"note":""})

ws.cell(row=r, column=1, value="")
ws.cell(row=r, column=2, value="顶部导航 — 出现在所有页面顶部")
ws.cell(row=r, column=2).font = SEC_FONT
ws.cell(row=r, column=2).fill = LIGHT_FILL
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
ws.row_dimensions[r].height = 18
r += 1
for i, row in enumerate(nav_rows2, 1):
    ws.cell(row=r, column=1, value=i).alignment = WRAP_C
    ws.cell(row=r, column=2, value=row["el"])
    ws.cell(row=r, column=3, value=row["txt"])
    ws.cell(row=r, column=5, value=row["note"])
    ws.cell(row=r, column=6, value=row["css"])
    for col in range(1,7):
        cell = ws.cell(row=r, column=col); cell.border = BORDER; cell.font = BODY_FONT
        if col in (2,3,5): cell.alignment = WRAP
    ws.cell(row=r, column=6).font = NOTE_FONT
    r += 1

# --- footers (two templates) ---
def footer_rows_of(fname):
    soup = BeautifulSoup(open(os.path.join(SRC, fname), encoding="utf-8").read(), "html.parser")
    for t in soup(STRIP): t.decompose()
    ft = soup.find("footer")
    out = []
    # brand
    brand = ft.select_one(".foot-brand p")
    if brand:
        out.append(("页脚标语(品牌区)","页脚标语 · 所有页面最下方标语", clean(brand.get_text(" ")), "footer .foot-brand p"))
    # columns
    for ci, col in enumerate(ft.select(".foot-top .foot-col"), 1):
        h4 = col.find("h4")
        coltitle = clean(h4.get_text(" ")) if h4 else "(栏目)"
        if coltitle:
            out.append((f"栏目[{coltitle}] · 栏目标题", "页脚栏目标题", coltitle, f"footer .foot-col:nth-of-type({ci}) h4"))
        for a in col.find_all("a"):
            t = clean(a.get_text(" "))
            if t:
                out.append((f"栏目[{coltitle}] · 链接文字", "页脚链接 · 栏目 " + coltitle, t, f"footer .foot-col:nth-of-type({ci}) a"))
        for p in col.find_all("p"):
            t = clean(p.get_text(" "))
            if t:
                out.append((f"栏目[{coltitle}] · 说明文字", "页脚文字 · 栏目 " + coltitle, t, f"footer .foot-col:nth-of-type({ci}) p"))
    bottom = ft.select_one(".foot-bottom")
    if bottom:
        spans = bottom.find_all("span", recursive=False)
        if spans:
            for si, sp in enumerate(spans, 1):
                t = clean(sp.get_text(" "))
                if t:
                    out.append(("页脚最底部 · 小字", "页脚最底行文字", t, f"footer .foot-bottom span:nth-of-type({si})"))
        else:
            # template B: text may sit in <p> or directly
            ps = bottom.find_all(["p","span","a"], recursive=False)
            if ps:
                for pi, p in enumerate(ps, 1):
                    t = clean(p.get_text(" "))
                    if t:
                        out.append(("页脚最底部 · 小字", "页脚最底行文字", t, f"footer .foot-bottom > :nth-of-type({pi})"))
            else:
                t = clean(bottom.get_text(" "))
                if t:
                    out.append(("页脚最底部 · 小字", "页脚最底行文字", t, "footer .foot-bottom"))
    return out

def write_footer_group(ws, r, group_title, pages_zh, items):
    ws.cell(row=r, column=2, value=group_title)
    ws.cell(row=r, column=2).font = SEC_FONT
    ws.cell(row=r, column=2).fill = YELL_FILL
    ws.cell(row=r, column=5, value="适用页面: " + pages_zh)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    ws.cell(row=r, column=5).font = NOTE_FONT
    ws.row_dimensions[r].height = 18
    r += 1
    # brand line first
    for item in items:
        pos, note, txt = item[0], item[1], item[2]
        css = item[3] if len(item) > 3 else ""
        ws.cell(row=r, column=1, value="")
        ws.cell(row=r, column=2, value=pos)
        ws.cell(row=r, column=3, value=txt)
        ws.cell(row=r, column=5, value=note)
        ws.cell(row=r, column=6, value=css)
        for col in range(1,7):
            cell = ws.cell(row=r, column=col); cell.border = BORDER; cell.font = BODY_FONT
            if col in (2,3,5): cell.alignment = WRAP
        ws.cell(row=r, column=6).font = NOTE_FONT
        r += 1
    return r

ws.cell(row=r, column=2, value="")
ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor="F2F2F2")
r += 1

r = write_footer_group(ws, r, "页脚 · 版本 A(精简,含微信/WhatsApp 与西语预告)",
    "首页 / 关于我们 / 合作大学(①②③)", footer_rows_of("index.html"))
r = write_footer_group(ws, r, "页脚 · 版本 B(完整,含栏目 Service lines 与 Building bridges 标语)",
    "中外合作办学 / 硕士专班 / 申请与录取 / 研学 / 联系我们(④⑤⑥⑦⑧)", footer_rows_of("cooperative-education.html"))
# note page-level footer differences inside version B
diff_rows = [
    ("栏目[Explore] · 链接文字", "Contact us", "仅 ⑦研学 与 ⑧联系 两页的 Explore 栏多出的链接;④⑤⑥ 页没有",
     "footer .foot-col:nth-of-type(2) a:last-of-type"),
]
for pos, txt, note, css in diff_rows:
    ws.cell(row=r, column=1, value="")
    ws.cell(row=r, column=2, value=pos)
    ws.cell(row=r, column=3, value=txt)
    ws.cell(row=r, column=5, value=note)
    ws.cell(row=r, column=6, value=css)
    for col in range(1,7):
        cell = ws.cell(row=r, column=col); cell.border = BORDER; cell.font = BODY_FONT
        if col in (2,3,5): cell.alignment = WRAP
    ws.cell(row=r, column=6).font = NOTE_FONT
    r += 1
ws.cell(row=r, column=2, value="⚠️ 提示:同一句页脚文字若同时出现在上面两套里(例如品牌标语、Copyright),需两处都改。建议直接告诉我“统一页脚”,我合并成一套并同步全站。")
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
ws.cell(row=r, column=2).font = Font(name="Microsoft YaHei", size=10, bold=True, color="C00000")
ws.cell(row=r, column=2).alignment = WRAP
ws.row_dimensions[r].height = 34

# ---------- Pages ----------
for fname, zh in PAGES:
    title, blocks = extract_page(fname)
    ws = wb.create_sheet(zh)
    set_widths(ws, [6, 26, 55, 40, 24, 66])
    headers = ["序号","网页位置","现在的英文文案","想改成什么(留空=不改)","备注","内部定位(勿改)"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)
    styled_sheet(ws, 1)
    ws.freeze_panes = "A2"
    r = 2
    # browser tab title
    ws.cell(row=r, column=2, value="浏览器标签页标题(SEO)")
    ws.cell(row=r, column=2).font = SEC_FONT
    ws.cell(row=r, column=2).fill = LIGHT_FILL
    ws.cell(row=r, column=3, value=title)
    ws.cell(row=r, column=3).alignment = WRAP
    ws.cell(row=r, column=5, value="显示在浏览器标签/搜索结果标题,可改")
    ws.cell(row=r, column=5).font = NOTE_FONT
    ws.cell(row=r, column=5).alignment = WRAP
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    ws.row_dimensions[r].height = 16
    for col in range(1,7):
        ws.cell(row=r, column=col).border = BORDER
        ws.cell(row=r, column=col).font = BODY_FONT
    r += 1
    # blocks
    for kind, label, rows in blocks:
        if kind in ("nav","footer"):
            continue  # common sheet already
        if kind == "hero":
            label = "横幅区(page hero) · " + (label[:46] + "…" if len(label) > 46 else label)
        else:
            label = "正文区 · " + (label[:46] + "…" if len(label) > 46 else label)
        r = write_block(ws, r, label, rows, kind)
    # note on common parts
    ws.cell(row=r, column=2, value="顶部导航与页脚见【1 全站通用】表(改一次全站生效)。")
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    ws.cell(row=r, column=2).font = NOTE_FONT
    ws.cell(row=r, column=2).alignment = WRAP

wb.save(OUT)
print("saved:", OUT)
